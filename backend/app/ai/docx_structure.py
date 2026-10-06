"""Structured DOCX extraction: blocks with numbering, styles and table coordinates.

WHY THIS EXISTS
`extraction.extract_docx` returns visible text only. For a contract that loses
what makes it citable: Word's automatic numbering is not in the text at all
("7.2" lives in numbering.xml), heading levels and styles disappear, and a
table becomes "cell | cell" lines with no row or column. The Document Analyzer
cites "Clause 7.2", so it needs the structure.

This module does NOT replace `extract_docx`, whose text contract other callers
(intake) rely on. It runs only when a caller asks for structure, in the same
isolated child process, after the same package validation
(`inspect_docx_package`: entry count, size, compression ratio, traversal) and
under the same block and character budgets.

WHAT IT RETURNS
`ExtractionResult` whose `text` is the CANONICAL TEXT, generated
deterministically from the blocks and nothing else, and whose `blocks` list
describes every block by offsets into that text:

    block_id      "b00001", stable for a given file (source order)
    order         source order
    type          heading | paragraph | list_item | table_cell
    style_id, style_name, outline_level (0-based, None for body text)
    start, end    the block's span in the canonical text (label included)
    text_start    where the block's own text begins (after the label)
    list          {num_id, abstract_id, level, format, label, resolved} | None
    table         {table_id, row, col, row_span, col_span, depth, parent_block_id} | None

Blocks carry NO document text; text is only ever cut from the canonical text.

NUMBERING (resolved, never invented)
  * paragraph numPr, else the style chain (basedOn), numId 0 = "no numbering";
  * num -> abstractNum, including numStyleLink indirection;
  * lvlOverride: startOverride and full level overrides;
  * counters kept per abstractNum (Word continues a list across num
    instances unless a num restarts it); deeper levels reset when a higher
    level advances, honouring lvlRestart; isLgl renders all levels decimal;
  * formats: decimal, decimalZero, lower/upperLetter, lower/upperRoman,
    ordinal, bullet (shown as "•"), none.
  Anything else -- an unknown num, a missing level, an unsupported format,
  LISTNUM/SEQ field numbering -- is NOT guessed: the block keeps no label,
  `list.resolved` is False and the limitation `numbering_unresolved` (or
  `numbering_field_unresolved`) is recorded.

TEXT RULES
  * Tracked changes: insertions are included, deletions excluded (the
    "accepted" reading) -> limitation `tracked_changes`.
  * Hidden text (w:vanish) is excluded, as a reader does not see it ->
    `hidden_text_excluded`. It is also where an injected instruction hides.
  * Content controls (w:sdt) at body, row, cell and run level are read.
  * Text boxes are not read -> `text_boxes_not_read`; embedded documents
    (altChunk) -> `embedded_content_not_read`.
  * Paragraph text is stripped; a numbered block renders "<label> <text>".
  * A table row renders as the legacy extractor renders it: each distinct
    cell's text (its paragraphs and nested tables, joined by spaces), the
    non-empty ones joined by " | ". Merged cells appear once, with spans.
  For a document with none of the constructs above (no numbering, sdt,
  tracked changes or hidden text), the canonical text equals
  `extract_docx(...).text` exactly -- a tested property.
"""
from __future__ import annotations

import re
from pathlib import Path

from app.ai.extraction import (
    COMPLETE,
    ERR_NO_TEXT_LAYER,
    ERR_PACKAGE_INVALID,
    ERR_PACKAGE_TOO_LARGE,
    ERR_PARSE_FAILED,
    EXTRACTOR_VERSION,
    MAX_DOCX_BLOCKS,
    MAX_EXTRACTED_CHARS,
    NONE,
    OUTCOME_SUCCEEDED,
    PARTIAL_OR_UNCERTAIN,
    ExtractionResult,
    _failed,
    _refused,
    inspect_docx_package,
)

#: Bump on any change to blocks, labels or canonical text. Travels on the
#: result as `extractor_version`, so an analysis records which reading of the
#: file it was built on. The shared EXTRACTOR_VERSION (intake) is untouched.
STRUCTURE_VERSION = "docx-blocks.1"
VERSION = f"{EXTRACTOR_VERSION}+{STRUCTURE_VERSION}"

MAX_TABLE_DEPTH = 10
MAX_STYLE_CHAIN = 20

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
RT_NUMBERING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"
RT_STYLES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"


def _q(tag: str) -> str:
    return f"{{{W}}}{tag}"


def _val(element, child: str | None = None, name: str = "val"):
    if element is None:
        return None
    if child is not None:
        element = element.find(_q(child))
        if element is None:
            return None
    return element.get(_q(name))


def _int(value, default=None):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _on(element) -> bool:
    """A Word on/off property: present and not explicitly false."""
    return element is not None and _val(element) not in ("0", "false", "off")


# ── styles ───────────────────────────────────────────────────────────────────

_HEADING_NAME = re.compile(r"^heading ([1-9])$", re.IGNORECASE)


class Styles:
    def __init__(self, root):
        self.by_id: dict[str, dict] = {}
        self.default_paragraph: str | None = None
        if root is None:
            return
        for s in root.iter(_q("style")):
            sid = s.get(_q("styleId"))
            if not sid:
                continue
            ppr = s.find(_q("pPr"))
            numpr = ppr.find(_q("numPr")) if ppr is not None else None
            self.by_id[sid] = {
                "type": s.get(_q("type")),
                "name": _val(s, "name"),
                "based_on": _val(s, "basedOn"),
                "num_id": _val(numpr, "numId") if numpr is not None else None,
                "ilvl": _int(_val(numpr, "ilvl")) if numpr is not None else None,
                "outline": _int(_val(ppr, "outlineLvl")) if ppr is not None else None,
            }
            if s.get(_q("type")) == "paragraph" and s.get(_q("default")) in ("1", "true"):
                self.default_paragraph = sid

    def chain(self, sid: str | None):
        seen = set()
        while sid and sid in self.by_id and sid not in seen and len(seen) < MAX_STYLE_CHAIN:
            seen.add(sid)
            yield sid, self.by_id[sid]
            sid = self.by_id[sid]["based_on"]

    def numbering(self, sid):
        for found, style in self.chain(sid):
            if style["num_id"] is not None:
                return style["num_id"], style["ilvl"], found
        return None

    def outline_level(self, sid):
        for _, style in self.chain(sid):
            if style["outline"] is not None:
                return style["outline"]
            match = _HEADING_NAME.match(style["name"] or "")
            if match:
                return int(match.group(1)) - 1
            if (style["name"] or "").lower() == "title":
                return 0
        return None

    def name(self, sid):
        return (self.by_id.get(sid) or {}).get("name")


# ── numbering ────────────────────────────────────────────────────────────────

SUPPORTED_FORMATS = {"decimal", "decimalZero", "lowerLetter", "upperLetter", "lowerRoman",
                     "upperRoman", "ordinal", "bullet", "none"}


def _level(lvl) -> dict:
    return {
        "start": _int(_val(lvl, "start"), 0),
        "format": _val(lvl, "numFmt") or "decimal",
        "text": _val(lvl, "lvlText"),
        "restart": _int(_val(lvl, "lvlRestart")),
        "is_lgl": _on(lvl.find(_q("isLgl"))),
        "pstyle": _val(lvl, "pStyle"),
    }


class Numbering:
    def __init__(self, root):
        self.abstracts: dict[str, dict] = {}
        self.nums: dict[str, dict] = {}
        if root is None:
            return
        for a in root.findall(_q("abstractNum")):
            aid = a.get(_q("abstractNumId"))
            levels = {}
            for lvl in a.findall(_q("lvl")):
                ilvl = _int(lvl.get(_q("ilvl")))
                if ilvl is not None:
                    levels[ilvl] = _level(lvl)
            self.abstracts[aid] = {"levels": levels, "num_style_link": _val(a, "numStyleLink")}
        for n in root.findall(_q("num")):
            nid = n.get(_q("numId"))
            overrides = {}
            for o in n.findall(_q("lvlOverride")):
                ilvl = _int(o.get(_q("ilvl")))
                if ilvl is None:
                    continue
                lvl = o.find(_q("lvl"))
                overrides[ilvl] = {"start": _int(_val(o, "startOverride")),
                                   "level": _level(lvl) if lvl is not None else None}
            self.nums[nid] = {"abstract": _val(n, "abstractNumId"), "overrides": overrides}

    def resolve(self, num_id: str, styles: Styles, depth: int = 0):
        """(abstract key, levels with this num's overrides applied) or None."""
        num = self.nums.get(num_id)
        if num is None or depth > 5:
            return None
        abstract = self.abstracts.get(num["abstract"])
        if abstract is None:
            return None
        key, levels = num["abstract"], dict(abstract["levels"])
        link = abstract["num_style_link"]
        if link:
            # The list's real definition lives behind a numbering style.
            linked = (styles.by_id.get(link) or {}).get("num_id")
            inner = self.resolve(linked, styles, depth + 1) if linked else None
            if inner is None:
                return None
            key, levels = inner[0], dict(inner[1])
        for ilvl, o in num["overrides"].items():
            if o["level"] is not None:
                levels[ilvl] = o["level"]
        return key, levels

    def level_for_style(self, num_id: str, style_id: str, styles: Styles):
        """The level linked (lvl/pStyle) to the paragraph's style, nearest first
        along its basedOn chain; level 0 when none is linked."""
        resolved = self.resolve(num_id, styles)
        if resolved:
            for sid, _ in styles.chain(style_id):
                for ilvl, level in resolved[1].items():
                    if level["pstyle"] == sid:
                        return ilvl
        return 0


_ROMAN = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
          (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))


def _roman(n: int) -> str | None:
    if not 0 < n < 4000:
        return None
    out = []
    for value, glyph in _ROMAN:
        while n >= value:
            out.append(glyph)
            n -= value
    return "".join(out)


def _letter(n: int) -> str | None:
    # Word's sequence: a..z, then aa..zz, aaa..: the letter repeats.
    if n < 1:
        return None
    return chr(ord("a") + (n - 1) % 26) * ((n - 1) // 26 + 1)


def _ordinal(n: int) -> str:
    suffix = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def format_number(n: int, fmt: str) -> str | None:
    if fmt == "decimal":
        return str(n)
    if fmt == "decimalZero":
        return f"{n:02d}" if 0 <= n < 10 else str(n)
    if fmt in ("lowerLetter", "upperLetter"):
        s = _letter(n)
        return s.upper() if s and fmt == "upperLetter" else s
    if fmt in ("lowerRoman", "upperRoman"):
        s = _roman(n)
        return s.upper() if s and fmt == "upperRoman" else s
    if fmt == "ordinal":
        return _ordinal(n) if n > 0 else None
    return None


class Counters:
    """Word's list counters: per abstract list, per level."""

    def __init__(self, numbering: Numbering, styles: Styles):
        self.numbering, self.styles = numbering, styles
        self.values: dict[str, dict[int, int]] = {}
        self.num_started: set[str] = set()

    def advance(self, num_id: str, ilvl: int):
        """The rendered label, or None when it cannot be resolved honestly."""
        resolved = self.numbering.resolve(num_id, self.styles)
        if resolved is None:
            return None, None, None
        key, levels = resolved
        level = levels.get(ilvl)
        if level is None:
            return key, None, None
        counters = self.values.setdefault(key, {})
        if num_id not in self.num_started:
            # A num that restarts its list (startOverride) takes effect the
            # first time it is used.
            self.num_started.add(num_id)
            for o_lvl, o in self.numbering.nums[num_id]["overrides"].items():
                if o["start"] is not None or o["level"] is not None:
                    counters.pop(o_lvl, None)
        start_of = {lv: (self.numbering.nums[num_id]["overrides"].get(lv, {}).get("start")
                         if self.numbering.nums[num_id]["overrides"].get(lv, {}).get("start") is not None
                         else spec["start"]) for lv, spec in levels.items()}
        counters[ilvl] = counters[ilvl] + 1 if ilvl in counters else start_of[ilvl]
        for deeper in [lv for lv in counters if lv > ilvl]:
            restart = (levels.get(deeper) or {}).get("restart")
            if restart == 0:
                continue                       # this level never restarts
            if restart is None or ilvl < restart:
                counters.pop(deeper)
        return key, level, self._render(level, levels, counters, start_of)

    @staticmethod
    def _render(level, levels, counters, start_of):
        fmt = level["format"]
        if fmt not in SUPPORTED_FORMATS:
            return None
        if fmt == "none":
            return ""
        if fmt == "bullet":
            return "•"
        template = level["text"]
        if template is None:
            return None
        failed = []

        def substitute(match):
            lv = int(match.group(1)) - 1
            spec = levels.get(lv)
            if spec is None:
                failed.append(lv)
                return ""
            value = counters.get(lv, start_of.get(lv, spec["start"]))
            shown = format_number(value, "decimal" if level["is_lgl"] else spec["format"])
            if shown is None:
                failed.append(lv)
                return ""
            return shown

        label = re.sub(r"%([1-9])", substitute, template)
        return None if failed else label.strip()


# ── text of one paragraph ────────────────────────────────────────────────────

_SKIP = {_q(t) for t in ("pPr", "rPr", "del", "moveFrom", "delText", "delInstrText", "instrText",
                         "txbxContent", "footnoteReference", "endnoteReference",
                         "commentReference", "fldData")}
_FIELD_NUMBERING = re.compile(r"\b(LISTNUM|SEQ|AUTONUM\w*)\b", re.IGNORECASE)


def _paragraph_text(p, flags: set[str]) -> str:
    out: list[str] = []

    def walk(element):
        for child in element:
            tag = child.tag
            if not isinstance(tag, str):
                continue                        # comments, processing instructions
            if tag in (_q("del"), _q("moveFrom")):
                flags.add("tracked_changes")
                continue
            if tag == _q("instrText") and _FIELD_NUMBERING.search(child.text or ""):
                flags.add("numbering_field_unresolved")
            if tag == _q("txbxContent"):
                flags.add("text_boxes_not_read")
                continue
            if tag in _SKIP:
                continue
            if tag in (_q("ins"), _q("moveTo")):
                flags.add("tracked_changes")
            if tag == _q("r"):
                rpr = child.find(_q("rPr"))
                if rpr is not None and (_on(rpr.find(_q("vanish"))) or _on(rpr.find(_q("specVanish")))):
                    flags.add("hidden_text_excluded")
                    continue
            if tag == _q("t"):
                out.append(child.text or "")
            elif tag in (_q("tab"), _q("ptab")):
                out.append("\t")
            elif tag == _q("cr") or (tag == _q("br") and child.get(_q("type"), "textWrapping")
                                     == "textWrapping"):
                out.append("\n")            # page/column breaks are not text (as python-docx)
            elif tag == _q("noBreakHyphen"):
                out.append("-")
            else:
                walk(child)

    walk(p)
    return "".join(out).strip()


# ── the document walk ────────────────────────────────────────────────────────

class _Builder:
    def __init__(self, numbering: Numbering, styles: Styles):
        self.styles = styles
        self.counters = Counters(numbering, styles)
        self.numbering = numbering
        self.flags: set[str] = set()
        self.budget = MAX_DOCX_BLOCKS
        self.order = 0
        self.tables = 0
        self.lines: list[tuple[str, list[dict]]] = []   # (line text, blocks with LINE-relative offsets)

    def _block(self, **fields) -> dict:
        self.order += 1
        return {"block_id": f"b{self.order:05d}", "order": self.order, **fields}

    # A paragraph as its own line, or None. Advances numbering even when empty.
    def paragraph(self, p) -> tuple[str, dict] | None:
        ppr = p.find(_q("pPr"))
        style_id = _val(ppr, "pStyle") if ppr is not None else None
        style_id = style_id or self.styles.default_paragraph
        text = _paragraph_text(p, self.flags)

        numpr = ppr.find(_q("numPr")) if ppr is not None else None
        num_id, ilvl = None, None
        if numpr is not None:
            num_id, ilvl = _val(numpr, "numId"), _int(_val(numpr, "ilvl"))
            if num_id is not None and ilvl is None:
                ilvl = 0
        if num_id is None:
            inherited = self.styles.numbering(style_id)
            if inherited:
                num_id, style_ilvl, _ = inherited
                if ilvl is None:                # a direct ilvl beats the style's
                    ilvl = (style_ilvl if style_ilvl is not None
                            else self.numbering.level_for_style(num_id, style_id, self.styles))
        listing = None
        label = None
        if num_id is not None and num_id != "0":
            key, level, label = self.counters.advance(num_id, ilvl or 0)
            resolved = label is not None
            if not resolved:
                self.flags.add("numbering_unresolved")
            listing = {"num_id": num_id, "abstract_id": key, "level": ilvl or 0,
                       "format": level["format"] if level else None,
                       "label": label if resolved else None, "resolved": resolved}

        if not text:
            return None
        outline = _int(_val(ppr, "outlineLvl")) if ppr is not None else None
        if outline is None:
            outline = self.styles.outline_level(style_id)
        if outline is not None and outline >= 9:
            outline = None                    # 9 = "body text"
        kind = "heading" if outline is not None else ("list_item" if listing else "paragraph")
        prefix = f"{label} " if label else ""
        line = prefix + text
        block = self._block(type=kind, style_id=style_id, style_name=self.styles.name(style_id),
                            outline_level=outline, start=0, end=len(line), text_start=len(prefix),
                            list=listing, table=None)
        return line, block

    def body(self, container):
        for child in self._children(container):
            if self.budget <= 0:
                self.flags.add("block_limit_reached")
                return
            tag = child.tag
            if tag == _q("p"):
                self.budget -= 1
                made = self.paragraph(child)
                if made:
                    line, block = made
                    self.lines.append((line, [block]))
            elif tag == _q("tbl"):
                self.budget -= 1
                for line, blocks in self.table(child, depth=0, parent=None):
                    self.lines.append((line, blocks))

    def _children(self, container):
        """Block-level children, looking through content controls and custom XML."""
        for child in container:
            tag = child.tag
            if not isinstance(tag, str):
                continue
            if tag == _q("sdt"):
                content = child.find(_q("sdtContent"))
                if content is not None:
                    yield from self._children(content)
            elif tag in (_q("customXml"), _q("smartTag")):
                yield from self._children(child)
            elif tag == _q("altChunk"):
                self.flags.add("embedded_content_not_read")
            elif tag in (_q("p"), _q("tbl"), _q("tr"), _q("tc")):
                yield child

    def table(self, tbl, depth: int, parent: str | None) -> list[tuple[str, list[dict]]]:
        """Each row as one line, rendered exactly like the legacy extractor."""
        if depth > MAX_TABLE_DEPTH:
            self.flags.add("nesting_limit_reached")
            return []
        self.tables += 1
        table_id = f"t{self.tables:04d}"
        lines = []
        open_merges: dict[int, dict] = {}         # grid column -> restart cell block
        row_index = -1
        for tr in self._children(tbl):
            if tr.tag != _q("tr"):
                continue
            row_index += 1
            trpr = tr.find(_q("trPr"))
            col = _int(_val(trpr, "gridBefore"), 0) if trpr is not None else 0
            cells: list[tuple[str, list[dict]]] = []
            for tc in self._children(tr):
                if tc.tag != _q("tc"):
                    continue
                if self.budget <= 0:
                    self.flags.add("block_limit_reached")
                    break
                self.budget -= 1
                tcpr = tc.find(_q("tcPr"))
                span = max(1, _int(_val(tcpr, "gridSpan"), 1)) if tcpr is not None else 1
                vmerge = tcpr.find(_q("vMerge")) if tcpr is not None else None
                if vmerge is not None and _val(vmerge) != "restart":
                    # Continuation of a vertically merged cell: the SAME cell,
                    # rendered once (legacy behaviour); its span grows.
                    head = open_merges.get(col)
                    if head is not None:
                        head["table"]["row_span"] += 1
                    col += span
                    continue
                block = self._block(type="table_cell", style_id=None, style_name=None,
                                    outline_level=None, start=0, end=0, text_start=0, list=None,
                                    table={"table_id": table_id, "row": row_index, "col": col,
                                           "row_span": 1, "col_span": span, "depth": depth,
                                           "parent_block_id": parent})
                if vmerge is not None:
                    open_merges[col] = block
                else:
                    open_merges.pop(col, None)
                cells.append(self._cell(tc, block, depth))
                col += span
            line, blocks = self._join(cells, " | ")
            stripped = line.strip(" |")
            if stripped != line:
                lead = len(line) - len(line.lstrip(" |"))
                line, blocks = stripped, [self._shift(b, -lead, len(stripped)) for b in blocks]
            if line:
                lines.append((line, blocks))
        return lines

    def _cell(self, tc, block: dict, depth: int) -> tuple[str, list[dict]]:
        """A cell's paragraphs and nested tables, joined by spaces (legacy)."""
        parts: list[tuple[str, list[dict]]] = []
        for child in self._children(tc):
            if child.tag == _q("p"):
                text = _paragraph_text(child, self.flags)
                if text:
                    parts.append((text, []))
            elif child.tag == _q("tbl"):
                parts.extend(self.table(child, depth + 1, block["block_id"]))
        text, inner = self._join(parts, " ")
        block["start"], block["end"], block["text_start"] = 0, len(text), 0
        return text, [block] + inner

    @staticmethod
    def _shift(block: dict, by: int, limit: int) -> dict:
        for key in ("start", "end", "text_start"):
            block[key] = min(max(block[key] + by, 0), limit)
        return block

    def _join(self, parts, sep: str) -> tuple[str, list[dict]]:
        """Join non-empty parts with `sep`; empty parts keep their blocks at the join point."""
        text, blocks = "", []
        for part, part_blocks in parts:
            if part and text:
                text += sep
            offset = len(text)
            blocks.extend(self._shift(b, offset, offset + len(part)) for b in part_blocks)
            text += part
        return text, blocks

    def canonical(self) -> tuple[str, list[dict]]:
        text, blocks = "", []
        for i, (line, line_blocks) in enumerate(self.lines):
            if i:
                text += "\n"
            offset = len(text)
            for b in line_blocks:
                blocks.append(self._shift(b, offset, offset + len(line)))
            text += line
        blocks.sort(key=lambda b: b["order"])
        return text, blocks


def _part_xml(document, reltype: str):
    try:
        return document.part.part_related_by(reltype).element
    except (KeyError, AttributeError):
        return None


def extract_docx_structured(path: Path) -> ExtractionResult:
    """Never raises: every failure is a coded result, like `extract_docx`."""
    error_code, unsupported = inspect_docx_package(path)
    if error_code:
        return (_refused(error_code) if error_code in (ERR_PACKAGE_INVALID, ERR_PACKAGE_TOO_LARGE)
                else _failed(error_code))
    try:
        import docx
        document = docx.Document(str(path))
        builder = _Builder(Numbering(_part_xml(document, RT_NUMBERING)),
                           Styles(_part_xml(document, RT_STYLES)))
        builder.body(document.element.body)
        text, blocks = builder.canonical()
    except Exception:
        return _failed(ERR_PARSE_FAILED)

    limitations = [f"unsupported_part:{name}" for name in unsupported] + sorted(builder.flags)
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS]
        blocks = [b for b in blocks if b["start"] < MAX_EXTRACTED_CHARS]
        for b in blocks:
            b["end"] = min(b["end"], MAX_EXTRACTED_CHARS)
            b["text_start"] = min(b["text_start"], b["end"])
        limitations.append("extracted_text_cap_reached")

    if not text.strip():
        result = ExtractionResult(outcome=OUTCOME_SUCCEEDED, completeness=NONE,
                                  error_code=ERR_NO_TEXT_LAYER, limitations=limitations,
                                  extractor_version=VERSION)
        result.blocks = []
        return result
    result = ExtractionResult(outcome=OUTCOME_SUCCEEDED,
                              completeness=PARTIAL_OR_UNCERTAIN if limitations else COMPLETE,
                              text=text, limitations=limitations, extractor_version=VERSION)
    result.blocks = blocks
    return result
