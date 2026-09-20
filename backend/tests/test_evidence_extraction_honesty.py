"""Partially read evidence must never look like fully read evidence.

THE DEFECT THIS FILE EXISTS FOR

A three-page bundle — a typed cover sheet in front of two scanned pages —
returned the cover sheet, no error, and status `readable`. The analysis then
reasoned about the case from a title page, and nothing in the record said that
two thirds of the evidence had not been seen. The same thing happened inside a
single page: a typed court header above a scanned body produced 39 characters
that were indistinguishable from a fully-read page.

Both are real client uploads, not edge cases. These tests build them.

WHAT IS DELIBERATELY NOT ASSERTED

That we can tell a scan from a blank page. We cannot — PDFs carry no structure
that says so — and any test asserting it would be pinning a guess. What is
asserted is that both are reported as *uncertain* rather than as read.

Fixtures are built here rather than committed: a real scanned FIR is client data,
and a synthetic one is honest about being synthetic.
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from app.ai import extraction as E

# Fixture construction needs a PDF writer and an image library. Both are already
# dependencies of the extraction path under test.
reportlab = pytest.importorskip("reportlab")
PIL = pytest.importorskip("PIL")


# ── fixture builders ────────────────────────────────────────────────────────

def _typed_pdf(path: Path, *pages: str) -> Path:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    for text in pages:
        c.drawString(72, 720, text)
        c.showPage()
    c.save()
    return path


def _image_pages(path: Path, count: int = 2) -> Path:
    from PIL import Image

    page = Image.new("RGB", (1200, 1600), "white")
    extra = [Image.new("RGB", (1200, 1600), "white") for _ in range(count - 1)]
    page.save(path, "PDF", save_all=bool(extra), append_images=extra)
    return path


def _concat(target: Path, *sources: Path) -> Path:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for src in sources:
        for page in PdfReader(str(src)).pages:
            writer.add_page(page)
    with target.open("wb") as handle:
        writer.write(handle)
    return target


def _header_over_scan(path: Path) -> Path:
    """One page: a real text layer above a full-page image."""
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    body = Image.new("RGB", (1400, 1700), "white")
    ImageDraw.Draw(body).rectangle([40, 40, 1360, 1660], outline="black", width=5)
    body_path = path.with_suffix(".png")
    body.save(body_path)

    c = canvas.Canvas(str(path), pagesize=letter)
    c.drawString(72, 760, "IN THE COURT OF THE CIVIL JUDGE, LAHORE")
    c.drawImage(ImageReader(str(body_path)), 72, 100, width=460, height=620)
    c.showPage()
    c.save()
    return path


def _docx(path: Path, build) -> Path:
    import docx

    document = docx.Document()
    build(document)
    document.save(str(path))
    return path


# ── the headline case ───────────────────────────────────────────────────────

def test_a_typed_cover_over_scans_is_not_reported_as_read(tmp_path):
    """The defect, exactly as it reached production."""
    bundle = _concat(
        tmp_path / "bundle.pdf",
        _typed_pdf(tmp_path / "cover.pdf", "COVER SHEET - Case No. 44 of 2026"),
        _image_pages(tmp_path / "scans.pdf", 2),
    )

    result = E.extract_file(str(bundle))

    assert result.outcome == E.OUTCOME_SUCCEEDED
    assert result.completeness == E.PARTIAL_OR_UNCERTAIN
    assert result.pages_total == 3
    assert result.pages_with_text == 1
    assert "COVER SHEET" in result.text


def test_a_header_above_a_scanned_body_is_uncertain_not_complete(tmp_path):
    """The same failure inside ONE page.

    The page yields text, so every page-level counter says it was read. Only the
    image reference says otherwise, which is why an image on a page is enough on
    its own to refuse the word `complete`.
    """
    result = E.extract_file(str(_header_over_scan(tmp_path / "hdr.pdf")))

    assert result.pages_with_text == 1
    assert result.pages_attempted == 1
    assert result.text_yielding_page_ratio == 1.0, (
        "every page yielded text — which is exactly why the ratio cannot be "
        "what decides completeness")
    assert result.completeness == E.PARTIAL_OR_UNCERTAIN
    assert result.page_reports[0].images_present is True


def test_a_fully_typed_pdf_is_complete(tmp_path):
    """The control. If this is not `complete`, the rule is useless."""
    result = E.extract_file(
        str(_typed_pdf(tmp_path / "typed.pdf", "Page one text.", "Page two text.")))

    assert result.completeness == E.COMPLETE
    assert result.pages_total == 2
    assert result.pages_with_text == 2
    assert result.processing_coverage == 1.0
    assert result.text_yielding_page_ratio == 1.0


def test_a_blank_page_is_not_called_a_scan(tmp_path):
    """A genuinely empty page and an image-only page give the same signal.

    Neither is labelled. Both are simply 'no text found'.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    blank = tmp_path / "blank.pdf"
    c = canvas.Canvas(str(blank), pagesize=letter)
    c.showPage()
    c.save()

    result = E.extract_file(str(blank))

    assert result.completeness == E.NONE
    assert result.error_code == E.ERR_NO_TEXT_LAYER
    assert result.page_reports[0].state == E.PAGE_NO_TEXT_FOUND
    assert "scan" not in result.page_reports[0].state


# ── the two ratios are distinct, and neither proves completeness ────────────

def test_the_two_ratios_measure_different_things(tmp_path):
    bundle = _concat(
        tmp_path / "b.pdf",
        _typed_pdf(tmp_path / "t.pdf", "one"),
        _image_pages(tmp_path / "s.pdf", 3),
    )
    result = E.extract_file(str(bundle))

    # Every page was looked at...
    assert result.processing_coverage == 1.0
    # ...and only a quarter of them gave anything up.
    assert result.text_yielding_page_ratio == pytest.approx(0.25)
    assert result.completeness == E.PARTIAL_OR_UNCERTAIN


def test_neither_ratio_being_one_is_enough_to_claim_complete(tmp_path):
    """Both ratios read 1.0 for the header-over-scan page. It is still uncertain."""
    result = E.extract_file(str(_header_over_scan(tmp_path / "h.pdf")))

    assert result.processing_coverage == 1.0
    assert result.text_yielding_page_ratio == 1.0
    assert result.completeness != E.COMPLETE


def test_a_page_cap_is_reported_as_skipped_not_absent(tmp_path):
    pages = [f"Typed page {n}." for n in range(6)]
    result = E.extract_pdf(_typed_pdf(tmp_path / "many.pdf", *pages), max_pages=2)

    assert result.pages_total == 6
    assert result.pages_attempted == 2
    assert result.pages_skipped == 4
    assert "page_limit_reached" in result.limitations
    assert result.processing_coverage == pytest.approx(2 / 6)
    assert result.completeness == E.PARTIAL_OR_UNCERTAIN


# ── DOCX: tables, nesting, merges, order ───────────────────────────────────

def test_docx_table_values_are_extracted(tmp_path):
    """The money and the date live in a table. They were being dropped."""
    def build(d):
        d.add_paragraph("Agreement between the parties.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Principal amount"
        t.cell(0, 1).text = "PKR 4,500,000"
        t.cell(1, 0).text = "Due date"
        t.cell(1, 1).text = "12 March 2026"

    result = E.extract_file(str(_docx(tmp_path / "t.docx", build)))

    assert "PKR 4,500,000" in result.text
    assert "12 March 2026" in result.text


def test_docx_document_order_is_preserved(tmp_path):
    """`.paragraphs` and `.tables` are two flat lists; order is not recoverable
    from them, which is why the body XML is walked instead."""
    def build(d):
        d.add_paragraph("FIRST")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "MIDDLE"
        d.add_paragraph("LAST")

    text = E.extract_file(str(_docx(tmp_path / "o.docx", build))).text

    assert text.index("FIRST") < text.index("MIDDLE") < text.index("LAST")


def test_docx_nested_tables_are_extracted(tmp_path):
    def build(d):
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "OUTER"
        t.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "INNER VALUE"

    text = E.extract_file(str(_docx(tmp_path / "n.docx", build))).text

    assert "INNER VALUE" in text


def test_a_merged_cell_is_not_duplicated(tmp_path):
    """`row.cells` returns one entry per grid position, so a merged cell comes
    back once per column it spans. A repeated amount is its own wrong answer."""
    def build(d):
        t = d.add_table(rows=2, cols=3)
        first = t.cell(0, 0)
        first.merge(t.cell(0, 1)).merge(t.cell(0, 2))
        first.text = "MERGED HEADER VALUE"

    text = E.extract_file(str(_docx(tmp_path / "m.docx", build))).text

    assert text.count("MERGED HEADER VALUE") == 1


def test_a_vertically_merged_cell_is_not_duplicated(tmp_path):
    def build(d):
        t = d.add_table(rows=3, cols=2)
        top = t.cell(0, 0)
        top.merge(t.cell(1, 0)).merge(t.cell(2, 0))
        top.text = "SPANS THREE ROWS"

    text = E.extract_file(str(_docx(tmp_path / "v.docx", build))).text

    assert text.count("SPANS THREE ROWS") == 1


def test_unread_docx_parts_downgrade_completeness(tmp_path):
    """A clause in a header is still a clause. Not reading it is disclosed."""
    def build(d):
        d.add_paragraph("Body text.")
        d.sections[0].header.paragraphs[0].text = "CONFIDENTIAL ANNEXURE"

    result = E.extract_file(str(_docx(tmp_path / "h.docx", build)))

    assert result.completeness == E.PARTIAL_OR_UNCERTAIN
    assert any(x.startswith("unsupported_part:") for x in result.limitations)


# ── package validation ──────────────────────────────────────────────────────

def test_an_ordinary_zip_is_refused_as_a_word_document(tmp_path):
    """`PK\\x03\\x04` is ZIP magic, not DOCX magic. Content-based upload
    validation accepted any archive and stored it as `.docx`."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("hello.txt", "an ordinary archive")
    forged = tmp_path / "forged.docx"
    forged.write_bytes(buf.getvalue())

    result = E.extract_file(str(forged))

    assert result.outcome == E.OUTCOME_REFUSED
    assert result.error_code == E.ERR_PACKAGE_INVALID


def test_a_zip_with_content_types_but_no_document_is_still_refused(tmp_path):
    """Checking for `[Content_Types].xml` alone is a shibboleth anyone can add."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("payload.bin", "x" * 100)
    forged = tmp_path / "half.docx"
    forged.write_bytes(buf.getvalue())

    assert E.extract_file(str(forged)).error_code == E.ERR_PACKAGE_INVALID


def test_an_over_expanding_package_is_refused_before_parsing(tmp_path):
    """Compressed upload size does not bound expanded size."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "A" * (12 * 1024 * 1024))
    bomb = tmp_path / "bomb.docx"
    bomb.write_bytes(buf.getvalue())

    result = E.extract_file(str(bomb))

    assert result.outcome == E.OUTCOME_REFUSED
    assert result.error_code == E.ERR_PACKAGE_TOO_LARGE
    assert bomb.stat().st_size < 1024 * 1024, (
        "fixture must be small on disk — that is the whole point")


def test_a_package_entry_escaping_its_directory_is_refused(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w/>")
        archive.writestr("../../evil.xml", "x")
    forged = tmp_path / "escape.docx"
    forged.write_bytes(buf.getvalue())

    assert E.extract_file(str(forged)).error_code == E.ERR_PACKAGE_INVALID


def test_a_real_docx_still_passes_validation(tmp_path):
    """The control for all of the above."""
    result = E.extract_file(
        str(_docx(tmp_path / "real.docx", lambda d: d.add_paragraph("Hello."))))

    assert result.outcome == E.OUTCOME_SUCCEEDED
    assert "Hello." in result.text


# ── legacy .doc and unsupported formats ─────────────────────────────────────

def test_a_legacy_doc_gets_its_own_actionable_message(tmp_path):
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16)

    result = E.extract_file(str(legacy))

    assert result.error_code == E.ERR_LEGACY_DOC_FORMAT
    message = E.message_for(result.error_code)
    assert ".docx" in message
    assert "still attached" in message, (
        "the user must be told their upload was not thrown away")


@pytest.mark.parametrize("name,magic", [
    ("photo.png", b"\x89PNG\r\n\x1a\n"),
    ("photo.jpg", b"\xff\xd8\xff"),
    ("photo.gif", b"GIF89a"),
    ("photo.webp", b"RIFF0000WEBP"),
])
def test_an_image_gets_its_own_truthful_refusal(tmp_path, name, magic):
    """Images are ACCEPTED at upload, so `unsupported_format` was a lie.

    A client who uploaded a photo of an FIR has not made the same mistake as one
    who uploaded a spreadsheet: the format is supported for storage, it is text
    recognition that does not exist. The remedies differ, so the codes do too.
    """
    image = tmp_path / name
    image.write_bytes(magic + b"\x00" * 32)

    result = E.extract_file(str(image))

    assert result.outcome == E.OUTCOME_REFUSED
    assert result.error_code == E.ERR_IMAGE_NO_TEXT_EXTRACTION

    message = E.message_for(result.error_code)
    assert "downloaded" in message, "the client must know the file was kept"
    assert "scan" not in message.lower(), "no claim about what the image contains"


def test_a_genuinely_unsupported_format_still_says_so(tmp_path):
    """The control: the image code must not swallow everything unreadable."""
    odd = tmp_path / "sheet.xlsx"
    odd.write_bytes(b"PK\x03\x04" + b"\x00" * 32)

    assert E.extract_file(str(odd)).error_code != E.ERR_IMAGE_NO_TEXT_EXTRACTION


def test_a_missing_file_is_distinguished_from_an_unreadable_one(tmp_path):
    assert E.extract_file(str(tmp_path / "nope.pdf")).error_code == E.ERR_FILE_MISSING


def test_a_malformed_pdf_fails_with_a_safe_code(tmp_path):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4\nthis is not a pdf body")

    result = E.extract_file(str(broken))

    assert result.outcome == E.OUTCOME_FAILED
    assert result.error_code in {E.ERR_PARSE_FAILED, E.ERR_FILE_UNREADABLE}


# ── error hygiene ───────────────────────────────────────────────────────────

def test_every_error_code_has_a_user_message():
    codes = [v for k, v in vars(E).items()
             if k.startswith("ERR_") and isinstance(v, str)]
    missing = [c for c in codes if c not in E.ERROR_MESSAGES]

    assert missing == [], f"error codes with no message: {missing}"


def test_error_messages_leak_no_paths_or_library_internals(tmp_path):
    """The old code returned `f"Could not read the file: {exc}"`, which put
    parser internals — and whatever they quoted — into model-visible text."""
    broken = tmp_path / "secret-client-name.pdf"
    broken.write_bytes(b"%PDF-1.4\ngarbage")

    result = E.extract_file(str(broken))
    message = E.message_for(result.error_code)

    assert "secret-client-name" not in message
    assert str(tmp_path) not in message
    for leak in ("Traceback", "pypdf", "docx", "zipfile", "Error(", "0x"):
        assert leak not in message


def test_no_error_message_mentions_a_file_path():
    for code, message in E.ERROR_MESSAGES.items():
        assert "/" not in message.replace("and/or", ""), code
        assert "\\" not in message, code
