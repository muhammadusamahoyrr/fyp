import sys
from pathlib import Path
from docx import Document

HERE = Path(__file__).resolve().parent
DOCX_PATH = HERE / "FIT_paper_v18.docx"

doc = Document(str(DOCX_PATH))

full_text = []

# Collect paragraph text
for p in doc.paragraphs:
    full_text.append(p.text)

# Collect table text
for t in doc.tables:
    for row in t.rows:
        row_str = " | ".join(cell.text.strip() for cell in row.cells)
        full_text.append(row_str)

all_str = "\n".join(full_text)

print("=== CHECKING REQUIRED STRINGS IN FIT_paper_v18.docx ===")

import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

required_strings = [
    "1,191 citation pairs",
    "531 improved",
    "123 worsened",
    "537 unchanged",
    "General model",
    "16.5",
    "0.0032",
    "largest improvement that remains",
    "10⁻⁵", # learning rate exponent
    "not specific to a single model configuration",
]

forbidden_strings = [
    "the three larger numbers did not",
    "each generalises past this system",
]

all_passed = True

for req in required_strings:
    if req in all_str:
        print(f"PASSED REQUIRED: '{req}'")
    else:
        print(f"FAILED MISSING REQUIRED: '{req}'")
        all_passed = False

for forb in forbidden_strings:
    if forb in all_str:
        print(f"FAILED FOUND FORBIDDEN: '{forb}'")
        all_passed = False
    else:
        print(f"PASSED FORBIDDEN ABSENT: '{forb}'")

# Inspect Table I in docx
print("\n=== TABLE I VISUAL CHECK IN DOCX ===")
for i, t in enumerate(doc.tables):
    for r in t.rows:
        row_cells = [c.text.strip() for c in r.cells]
        if "Property" in row_cells or "Training pairs" in row_cells:
            print("Table I row:", row_cells)

if all_passed:
    print("\n🎉 ALL DOCX VERIFICATION CHECKS PASSED PERFECTLY!")
else:
    print("\n⚠️ SOME CHECKS FAILED!")
    sys.exit(1)
